"""Что внутри фото или документа: текст и картинки для модели (docs/media.md).

Чистые синхронные функции без сети и базы; работа процессорная, вызывающий запускает их
в отдельном потоке. Наружу выходят только два исключения: `Unsupported` (такой вид файла
не разбираем) и `BadFile` (файл повреждён, защищён паролем или слишком велик для разбора);
текст обоих — короткая причина по-русски, её можно показать владельцу.

Виды файлов (`sniff`): сначала сигнатура содержимого, затем mime и расширение имени.
  image — JPEG, PNG, WebP, GIF (у анимации — первый кадр);
  pdf   — текст страниц, у скана — ещё и картинки первых страниц;
  docx  — абзацы и таблицы основного текста документа Word;
  xlsx  — значения ячеек по листам Excel;
  text  — txt, csv, md в UTF-8, UTF-16 с BOM или cp1251.
Не разбираем: старые .doc и .xls, презентации, архивы, HEIC, видео, звук и прочее —
`sniff` возвращает None, `extract` бросает `Unsupported`. Имя и mime приходят от
отправителя, поэтому сигнатура известного неподдерживаемого вида главнее расширения.

Пределы — от времени и памяти, файл приходит от чужих людей:
  фото — не больше MAX_PIXELS (50 Мп) по заголовку, до раскодирования; иначе BadFile;
         результат — JPEG с длинной стороной не больше image_side;
  PDF  — текст первых max_pages страниц и не больше max_text символов; картинки —
         первые render_pages страниц, разрешение до RENDER_DPI и длинная сторона не больше
         image_side; PDFium не потокобезопасен, поэтому весь разбор PDF идёт под одной
         блокировкой процесса (`_PDFIUM_LOCK`);
  docx, xlsx — архив ZIP: не больше ZIP_MAX_MEMBERS файлов внутри, заявленный распакованный
         объём не больше ZIP_MAX_TOTAL (100 МБ), степень сжатия большого файла внутри не
         больше ZIP_MAX_RATIO; XML с DOCTYPE или сущностями отклоняется; xlsx — не больше
         XLSX_MAX_SHEETS листов, XLSX_MAX_ROWS строк и XLSX_MAX_COLS столбцов на лист;
  текст — не больше max_text символов.
Время разбора одной страницы PDF самой библиотекой не ограничено: страница с огромным
числом объектов может считаться долго. Предел — число страниц и пикселей.

Флаг `truncated` — сработал какой-нибудь предел: текст обрезан, страницы PDF или строки
таблицы прочитаны не все, у скана показаны не все страницы.
"""

from __future__ import annotations

import io
import re
import threading
import zipfile
from dataclasses import dataclass, field
from xml.etree import ElementTree as ET

KINDS = ("image", "pdf", "docx", "xlsx", "text")

MAX_PIXELS = 50_000_000          # фото и кадр GIF: ширина × высота по заголовку
JPEG_QUALITY = 85
RENDER_DPI = 130                 # страница скана для модели
SCAN_CHARS_PER_PAGE = 50         # букв и цифр в среднем на страницу — меньше значит скан
ZIP_MAX_TOTAL = 100 * 1024 * 1024
ZIP_MAX_MEMBERS = 10_000
ZIP_MAX_RATIO = 200              # распакованный / сжатый для файла внутри больше ZIP_RATIO_FROM
ZIP_RATIO_FROM = 1024 * 1024
XLSX_MAX_SHEETS = 100
XLSX_MAX_ROWS = 2000
XLSX_MAX_COLS = 50

_PDFIUM_LOCK = threading.Lock()


class Unsupported(Exception):
    """Вид файла не разбираем; str(exc) — короткая причина по-русски."""


class BadFile(Exception):
    """Файл повреждён, защищён паролем или слишком сложен; str(exc) — причина по-русски."""


@dataclass
class Extracted:
    kind: str                    # "image" | "pdf" | "docx" | "xlsx" | "text"
    text: str                    # извлечённый текст ('' у картинки и скана без текста)
    pages: int | None            # страниц PDF, листов XLSX; иначе None
    images: list[bytes] = field(default_factory=list)  # JPEG для показа модели
    truncated: bool = False      # текст или страницы обрезаны пределами


# --- вид файла ---------------------------------------------------------------------------

_MIME = {
    "image/jpeg": "image", "image/jpg": "image", "image/pjpeg": "image",
    "image/png": "image", "image/webp": "image", "image/gif": "image",
    "application/pdf": "pdf", "application/x-pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/plain": "text", "text/csv": "text", "text/markdown": "text",
    "text/x-markdown": "text",
}
_EXT = {
    "jpg": "image", "jpeg": "image", "png": "image", "webp": "image", "gif": "image",
    "pdf": "pdf", "docx": "docx", "xlsx": "xlsx",
    "txt": "text", "csv": "text", "md": "text", "markdown": "text",
}

_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ENCRYPTED_PACKAGE = "EncryptedPackage".encode("utf-16-le")
_NOT_TEXT_PREFIXES = (
    # известные двоичные виды, которые не разбираем: главнее расширения и mime
    (b"Rar!\x1a\x07", "архивы не разбираем"),
    (b"7z\xbc\xaf\x27\x1c", "архивы не разбираем"),
    (b"\x1f\x8b", "архивы не разбираем"),
    (b"BZh", "архивы не разбираем"),
    (b"\xfd7zXZ\x00", "архивы не разбираем"),
    (b"OggS", "звук и видео не разбираем"),
    (b"ID3", "звук и видео не разбираем"),
    (b"fLaC", "звук и видео не разбираем"),
    (b"\x1aE\xdf\xa3", "звук и видео не разбираем"),
    (b"II*\x00", "картинки TIFF не разбираем"),
    (b"MM\x00*", "картинки TIFF не разбираем"),
    (b"BM", None),               # BMP — только вместе с расширением, см. _signature
    (b"\x7fELF", "программы не разбираем"),
    (b"MZ", None),
)


def _ext(name: str | None) -> str:
    if not name or "." not in name:
        return ""
    return name.rsplit(".", 1)[1].strip().lower()


def _mime(mime: str | None) -> str:
    return (mime or "").split(";", 1)[0].strip().lower()


def _hint(name: str | None, mime: str | None) -> str | None:
    """Вид по mime, затем по расширению — то, что заявил отправитель."""
    return _MIME.get(_mime(mime)) or _EXT.get(_ext(name))


def _zip_kind(data: bytes) -> tuple[str | None, str | None]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = set(zf.namelist())
    except Exception:  # noqa: BLE001 — битый архив: решают расширение и mime
        return "", None
    if "word/document.xml" in names:
        return "docx", None
    if "xl/workbook.xml" in names:
        return "xlsx", None
    if any(n.startswith("ppt/") for n in names):
        return None, "презентации не разбираем"
    return None, "архивы не разбираем"


def _signature(data: bytes, name: str | None, mime: str | None) -> tuple[str | None, str | None]:
    """(вид, причина отказа). Вид "" — сигнатура ничего не сказала, решает заявленное."""
    head = data[:64]
    if head.startswith(b"\xff\xd8\xff") or head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image", None
    if head[:6] in (b"GIF87a", b"GIF89a"):
        return "image", None
    if head[:4] == b"RIFF" and head[8:12] == b"WEBP":
        return "image", None
    if head[:4] == b"RIFF":
        return None, "звук и видео не разбираем"
    if b"%PDF-" in data[:1024]:
        return "pdf", None
    if head.startswith(b"PK\x03\x04") or head.startswith(b"PK\x05\x06"):
        kind, reason = _zip_kind(data)
        if kind == "":
            hint = _hint(name, mime)
            return (hint, None) if hint in ("docx", "xlsx") else (None, "архив повреждён или не разбираем")
        return kind, reason
    if head.startswith(_OLE):
        # Зашифрованный docx или xlsx — это файл OLE с потоком EncryptedPackage.
        hint = _hint(name, mime)
        if _ENCRYPTED_PACKAGE in data and hint in ("docx", "xlsx"):
            return hint, None
        return None, "старые форматы Word, Excel и PowerPoint (.doc, .xls, .ppt) не разбираем"
    if head[4:8] == b"ftyp":
        brand = head[8:12]
        if brand in (b"heic", b"heix", b"hevc", b"heim", b"heis", b"mif1", b"msf1"):
            return None, "фото HEIC не разбираем"
        if brand in (b"avif", b"avis"):
            return None, "фото AVIF не разбираем"
        return None, "звук и видео не разбираем"
    if head.startswith((b"\xef\xbb\xbf", b"\xff\xfe", b"\xfe\xff")):
        return "text", None
    for prefix, reason in _NOT_TEXT_PREFIXES:
        if head.startswith(prefix):
            if reason is None:
                if _hint(name, mime) == "text":
                    return "", None
                return None, "такой вид файла не разбираем"
            return None, reason
    return "", None


def _detect(data: bytes, name: str | None, mime: str | None) -> tuple[str | None, str]:
    if not data:
        return None, "файл пустой"
    kind, reason = _signature(data, name, mime)
    if kind:
        return kind, ""
    if kind is None:
        return None, reason or "такой вид файла не разбираем"
    hint = _hint(name, mime)
    if hint == "text":
        return "text", ""
    if hint is None:
        return None, "такой вид файла не разбираем"
    # Заявлен PDF, картинка или документ Office, а сигнатуры нет — файл не того вида.
    return hint, ""


def sniff(data: bytes, name: str | None, mime: str | None) -> str | None:
    """Вид файла: "image" | "pdf" | "docx" | "xlsx" | "text"; None — не разбираем."""
    return _detect(data, name, mime)[0]


# --- общие помощники ---------------------------------------------------------------------

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f\ufeff\ufffe\uffff]")
_SPACES = re.compile(r"[ \u00a0]+\n")
_BLANKS = re.compile(r"\n{3,}")


def _clean(text: str) -> str:
    """Убирает NUL и управляющие символы, лишние пустые строки и пробелы в концах строк."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL.sub("", text)
    text = _SPACES.sub("\n", text)
    text = _BLANKS.sub("\n\n", text)
    return text.strip()


class _Text:
    """Набирает текст до предела символов и помнит, что предел сработал."""

    def __init__(self, limit: int) -> None:
        self.limit = max(0, int(limit))
        self.parts: list[str] = []
        self.size = 0
        self.truncated = False

    @property
    def full(self) -> bool:
        return self.size >= self.limit

    def add(self, piece: str) -> bool:
        """Добавляет кусок; False — предел достигнут, дальше читать незачем."""
        if self.full:
            self.truncated = True
            return False
        room = self.limit - self.size
        if len(piece) > room:
            piece = piece[:room]
            self.truncated = True
        self.parts.append(piece)
        self.size += len(piece)
        return not self.truncated

    def text(self) -> str:
        return _clean("".join(self.parts))


def _jpeg(img) -> bytes:  # noqa: ANN001 — PIL.Image.Image
    out = io.BytesIO()
    img.save(out, format="JPEG", quality=JPEG_QUALITY, optimize=True)
    return out.getvalue()


def _rgb(img):  # noqa: ANN001, ANN202
    """RGB без прозрачности: прозрачное — на белом фоне."""
    from PIL import Image

    if img.mode.startswith("I;16") or img.mode == "I":
        # 16-битное серое: простое convert("RGB") обрезает всё ярче 255 в белое.
        img = img.convert("I").point(lambda v: v * (1 / 256)).convert("L")
    if img.mode in ("RGBA", "LA", "PA") or (img.mode == "P" and "transparency" in img.info):
        img = img.convert("RGBA")
        bg = Image.new("RGB", img.size, (255, 255, 255))
        bg.paste(img, mask=img.getchannel("A"))
        return bg
    if img.mode != "RGB":
        return img.convert("RGB")
    return img


# --- фото --------------------------------------------------------------------------------

def _image(data: bytes, image_side: int) -> Extracted:
    from PIL import Image, ImageOps

    side = max(16, int(image_side))
    try:
        with Image.open(io.BytesIO(data), formats=["JPEG", "PNG", "WEBP", "GIF"]) as img:
            w, h = img.size
            if w <= 0 or h <= 0:
                raise BadFile("картинка повреждена")
            if w * h > MAX_PIXELS:
                raise BadFile(f"картинка слишком большая: {w}×{h} точек")
            if getattr(img, "is_animated", False):
                img.seek(0)
            if img.format == "JPEG":
                img.draft("RGB", (side, side))  # раскодирует сразу в уменьшенном виде
            img.load()
            try:
                frame = ImageOps.exif_transpose(img)
            except Exception:  # noqa: BLE001 — битые сведения EXIF: оставляем как есть
                frame = img
            frame = _rgb(frame)
            frame.thumbnail((side, side), Image.Resampling.LANCZOS)
            return Extracted(kind="image", text="", pages=None, images=[_jpeg(frame)])
    except (BadFile, Unsupported):
        raise
    except Image.DecompressionBombError as exc:
        raise BadFile("картинка слишком большая") from exc
    except MemoryError as exc:
        raise BadFile("картинке не хватило памяти") from exc
    except Exception as exc:  # noqa: BLE001 — Pillow бросает разное на битых файлах
        raise BadFile("картинка повреждена или не читается") from exc


# --- PDF ---------------------------------------------------------------------------------

def _pdf(data: bytes, max_text: int, max_pages: int, render_pages: int, image_side: int) -> Extracted:
    try:
        import pypdfium2 as pdfium
    except ImportError as exc:  # pragma: no cover — зависимость образа
        raise Unsupported("разбор PDF недоступен") from exc

    with _PDFIUM_LOCK:
        try:
            doc = pdfium.PdfDocument(data)
        except pdfium.PdfiumError as exc:
            if getattr(exc, "err_code", None) == 4:  # FPDF_ERR_PASSWORD
                raise BadFile("файл защищён паролем") from exc
            if getattr(exc, "err_code", None) == 5:  # FPDF_ERR_SECURITY
                raise BadFile("файл зашифрован неизвестным способом") from exc
            raise BadFile("PDF повреждён") from exc
        except MemoryError as exc:
            raise BadFile("PDF не хватило памяти") from exc
        except Exception as exc:  # noqa: BLE001
            raise BadFile("PDF повреждён") from exc
        try:
            return _pdf_doc(doc, max_text, max_pages, render_pages, image_side)
        except (BadFile, Unsupported):
            raise
        except MemoryError as exc:
            raise BadFile("PDF не хватило памяти") from exc
        except Exception as exc:  # noqa: BLE001
            raise BadFile("PDF не читается") from exc
        finally:
            doc.close()


def _pdf_doc(doc, max_text: int, max_pages: int, render_pages: int, image_side: int) -> Extracted:  # noqa: ANN001
    total = len(doc)
    if total <= 0:
        raise BadFile("в PDF нет страниц")
    acc = _Text(max_text)
    read = 0
    meaningful = 0
    truncated = total > max_pages
    for index in range(min(total, max(0, int(max_pages)))):
        page = doc[index]
        try:
            textpage = page.get_textpage()
            try:
                piece = textpage.get_text_bounded()
            finally:
                textpage.close()
        finally:
            page.close()
        read += 1
        meaningful += sum(1 for ch in piece if ch.isalnum())
        if not acc.add(piece.strip() + "\n\n"):
            break
    truncated = truncated or acc.truncated
    images: list[bytes] = []
    if read and meaningful / read < SCAN_CHARS_PER_PAGE and not acc.truncated:
        count = min(total, max(0, int(render_pages)))
        for index in range(count):
            images.append(_render(doc, index, image_side))
        truncated = truncated or count < total
    return Extracted(kind="pdf", text=acc.text(), pages=total, images=images, truncated=truncated)


def _render(doc, index: int, image_side: int) -> bytes:  # noqa: ANN001
    page = doc[index]
    try:
        w, h = page.get_size()
        longest = max(w, h, 1.0)
        scale = min(RENDER_DPI / 72, max(16, int(image_side)) / longest)
        bitmap = page.render(scale=scale)
        try:
            img = bitmap.to_pil()
            return _jpeg(_rgb(img))
        finally:
            bitmap.close()
    finally:
        page.close()


# --- ZIP (docx, xlsx) --------------------------------------------------------------------

def _office_zip(data: bytes, kind: str) -> zipfile.ZipFile:
    """Открывает docx или xlsx с проверкой объёма и вида; иначе BadFile или Unsupported."""
    if data.startswith(_OLE):
        if _ENCRYPTED_PACKAGE in data:
            raise BadFile("файл защищён паролем")
        raise Unsupported("старые форматы Word и Excel (.doc, .xls) не разбираем")
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
        infos = zf.infolist()
    except Exception as exc:  # noqa: BLE001
        raise BadFile("файл повреждён") from exc
    try:
        if len(infos) > ZIP_MAX_MEMBERS:
            raise BadFile("в файле слишком много частей")
        total = 0
        for info in infos:
            if info.flag_bits & 0x1:
                raise BadFile("файл защищён паролем")
            total += info.file_size
            if total > ZIP_MAX_TOTAL:
                raise BadFile("файл слишком большой в распакованном виде")
            if info.file_size > ZIP_RATIO_FROM and info.file_size > ZIP_MAX_RATIO * max(info.compress_size, 1):
                raise BadFile("файл подозрительно сильно сжат")
        names = {i.filename for i in infos}
        need = "word/document.xml" if kind == "docx" else "xl/workbook.xml"
        if need not in names:
            raise BadFile("файл повреждён: нет основной части документа")
    except BaseException:
        zf.close()
        raise
    return zf


def _read_member(zf: zipfile.ZipFile, name: str) -> bytes:
    try:
        with zf.open(name) as fp:
            raw = fp.read(ZIP_MAX_TOTAL + 1)
    except Exception as exc:  # noqa: BLE001 — битое сжатие, неверная контрольная сумма
        raise BadFile("файл повреждён") from exc
    if len(raw) > ZIP_MAX_TOTAL:
        raise BadFile("файл слишком большой в распакованном виде")
    _no_doctype(raw)
    return raw


def _no_doctype(raw: bytes) -> None:
    """В OOXML не бывает DOCTYPE и сущностей; встретились — файл собран нарочно."""
    head = raw[:4096]
    if b"<!DOCTYPE" in head or b"<!ENTITY" in raw:
        raise BadFile("файл содержит недопустимую разметку")


# --- docx --------------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_W_P, _W_TBL, _W_TR, _W_TC = _W + "p", _W + "tbl", _W + "tr", _W + "tc"
_W_T, _W_TAB, _W_BR, _W_CR = _W + "t", _W + "tab", _W + "br", _W + "cr"
_W_SKIP = {_W + "delText", _W + "instrText"}
_W_CONTAINERS = {_W + "sdt", _W + "sdtContent", _W + "customXml", _W + "smartTag",
                 _W + "ins", _W + "moveTo"}


def _para_text(p: ET.Element) -> str:
    out: list[str] = []
    for el in p.iter():
        tag = el.tag
        if tag == _W_T and el.text:
            out.append(el.text)
        elif tag == _W_TAB:
            out.append("\t")
        elif tag in (_W_BR, _W_CR):
            out.append("\n")
    return "".join(out)


def _cell_text(tc: ET.Element) -> str:
    parts = [_para_text(p).strip() for p in tc.iter(_W_P)]
    return " ".join(p for p in parts if p).replace("\n", " ")


def _blocks(parent: ET.Element):  # noqa: ANN202 — генератор строк
    for child in parent:
        tag = child.tag
        if tag == _W_P:
            yield _para_text(child)
        elif tag == _W_TBL:
            for tr in child.iter(_W_TR):
                cells = [_cell_text(tc) for tc in tr if tc.tag == _W_TC]
                if any(cells):
                    yield " | ".join(cells)
            yield ""
        elif tag in _W_CONTAINERS:
            yield from _blocks(child)


def _docx(data: bytes, max_text: int) -> Extracted:
    zf = _office_zip(data, "docx")
    try:
        raw = _read_member(zf, "word/document.xml")
    finally:
        zf.close()
    try:
        root = ET.fromstring(raw)
    except MemoryError as exc:
        raise BadFile("документу не хватило памяти") from exc
    except Exception as exc:  # noqa: BLE001 — ParseError и прочее от разбора XML
        raise BadFile("документ повреждён") from exc
    body = root.find(_W + "body")
    acc = _Text(max_text)
    if body is not None:
        for line in _blocks(body):
            if not acc.add(line + "\n"):
                break
    return Extracted(kind="docx", text=acc.text(), pages=None, images=[], truncated=acc.truncated)


# --- xlsx --------------------------------------------------------------------------------

def _cell(value) -> str:  # noqa: ANN001
    if value is None:
        return ""
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if hasattr(value, "isoformat"):
        text = value.isoformat()
        return text[:-9] if text.endswith("T00:00:00") else text
    return str(value).replace("\r", " ").replace("\n", " ").replace("\t", " ").strip()


def _xlsx(data: bytes, max_text: int) -> Extracted:
    zf = _office_zip(data, "xlsx")
    try:
        for info in zf.infolist():
            if info.filename.endswith((".xml", ".rels")):
                _read_member(zf, info.filename)
    finally:
        zf.close()
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover — зависимость образа
        raise Unsupported("разбор таблиц недоступен") from exc
    import warnings

    acc = _Text(max_text)
    truncated = False
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True,
                                        keep_links=False)
            try:
                sheets = wb.worksheets
                pages = len(wb.sheetnames)
                if pages > XLSX_MAX_SHEETS:
                    truncated = True
                for ws in sheets[:XLSX_MAX_SHEETS]:
                    if not hasattr(ws, "iter_rows"):
                        continue  # лист-диаграмма
                    if not acc.add(f"Лист: {ws.title}\n"):
                        break
                    seen = 0
                    for row in ws.iter_rows(max_row=XLSX_MAX_ROWS + 1,
                                            max_col=XLSX_MAX_COLS + 1, values_only=True):
                        seen += 1
                        if seen > XLSX_MAX_ROWS:
                            truncated = True
                            break
                        if len(row) > XLSX_MAX_COLS and row[XLSX_MAX_COLS] is not None:
                            truncated = True
                        values = [_cell(v) for v in row[:XLSX_MAX_COLS]]
                        values = [v for v in values if v]
                        if values and not acc.add("\t".join(values) + "\n"):
                            break
                    if acc.full:
                        break
                    acc.add("\n")
            finally:
                wb.close()
    except (BadFile, Unsupported):
        raise
    except MemoryError as exc:
        raise BadFile("таблице не хватило памяти") from exc
    except Exception as exc:  # noqa: BLE001 — openpyxl бросает разное на битых файлах
        raise BadFile("таблица повреждена") from exc
    return Extracted(kind="xlsx", text=acc.text(), pages=pages, images=[],
                     truncated=truncated or acc.truncated)


# --- текст -------------------------------------------------------------------------------

def _decode(data: bytes, cut: bool) -> str:
    """UTF-8 или UTF-16 с BOM, UTF-8, иначе cp1251. cut — данные обрезаны по пределу,
    последний символ UTF-8 мог разорваться."""
    if data.startswith(b"\xef\xbb\xbf"):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        if cut and len(data) % 2:
            data = data[:-1]
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        if cut and exc.start >= len(data) - 3 and exc.reason == "unexpected end of data":
            return data[: exc.start].decode("utf-8")
    return data.decode("cp1251", errors="replace")


def _plain(data: bytes, max_text: int) -> Extracted:
    # Больше, чем нужно для max_text, не раскодируем: в UTF-8 символ — до 4 байт.
    limit = max(0, int(max_text))
    chunk = data[: limit * 4 + 4]
    cut = len(chunk) < len(data)
    text = _decode(chunk, cut)
    sample = text[:4096]
    if sample:
        bad = sum(1 for ch in sample if ch < " " and ch not in "\t\n\r\f")
        if bad / len(sample) > 0.05:
            raise BadFile("файл не похож на текст")
    acc = _Text(limit)
    acc.add(text)
    truncated = acc.truncated or cut
    return Extracted(kind="text", text=acc.text(), pages=None, images=[], truncated=truncated)


# --- вход --------------------------------------------------------------------------------

def extract(data: bytes, name: str | None, mime: str | None, *,
            max_text: int = 40_000, max_pages: int = 300, render_pages: int = 4,
            image_side: int = 1600) -> Extracted:
    """Разбирает файл. Unsupported — вид не разбираем; BadFile — файл не прочитать."""
    if not isinstance(data, (bytes, bytearray, memoryview)):
        raise BadFile("нет содержимого файла")
    data = bytes(data)
    kind, reason = _detect(data, name, mime)
    if kind is None:
        if reason == "файл пустой":
            raise BadFile(reason)
        raise Unsupported(reason)
    try:
        if kind == "image":
            return _image(data, image_side)
        if kind == "pdf":
            return _pdf(data, max_text, max_pages, render_pages, image_side)
        if kind == "docx":
            return _docx(data, max_text)
        if kind == "xlsx":
            return _xlsx(data, max_text)
        return _plain(data, max_text)
    except (BadFile, Unsupported):
        raise
    except MemoryError as exc:
        raise BadFile("файлу не хватило памяти") from exc
    except RecursionError as exc:
        raise BadFile("файл слишком сложен") from exc
    except Exception as exc:  # noqa: BLE001 — последний рубеж: наружу только два вида
        raise BadFile("файл не читается") from exc
